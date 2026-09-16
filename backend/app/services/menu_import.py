from __future__ import annotations

import csv
import io
import re
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable

from openpyxl import load_workbook


OPERATIONAL_COLUMNS = {
    "KIT - DESCARTAVEIS",
    "KIT - QUIMICO",
    "KIT - TEMPERO",
    "KIT - GALETEIRO",
}

CATEGORY_MAP = {
    "PRATO PRINCIPAL": "prato_principal",
    "PRATO PRINCIPAL 2": "prato_principal",
    "OPCAO AO PP": "opcao_prato_principal",
    "GUARNICAO": "guarnicao",
    "SALADA": "salada",
    "SALADA 2": "salada",
    "SALADA 3": "salada",
    "SOBREMESA": "sobremesa",
    "SOBREMESA 2": "sobremesa",
    "ARROZ": "arroz",
    "ARROZ 2": "arroz",
    "FEIJAO": "feijao",
    "BEBIDA": "bebida",
    "BEBIDA 2": "bebida",
    "BEBIDA 3": "bebida",
    "ACOMPANHAMENTO": "acompanhamento",
}

_COST_RE = re.compile(r"\s+-\s+(\d+(?:[.,]\d+)?)\s*$")
_PORTION_RE = re.compile(r"\(([^)]+)\)")
_FULL_CODE_RE = re.compile(r"\b\d{2}\.\d{2}\.\d{2}\.\d{3}(?:-\d+)?\b")
_PERCENT_PREFIX_RE = re.compile(r"^\s*\d+%\s*-\s*")


@dataclass(frozen=True)
class ImportedMenuItem:
    day: int
    source_column: str
    category: str
    name: str
    portion: str | None
    technical_sheet_code: str | None
    raw_value: str


def _normalize_header(value: object) -> str:
    return str(value or "").strip().upper()


def parse_planning_cell(raw: object) -> tuple[str, str | None, str | None] | None:
    if raw is None:
        return None
    text = str(raw).strip()
    if not text:
        return None

    # Custo per capita é dado operacional/comercial e morre na leitura.
    text_without_cost = _COST_RE.sub("", text).strip()

    portion_match = _PORTION_RE.search(text_without_cost)
    portion = portion_match.group(1).strip() if portion_match else None

    code_match = _FULL_CODE_RE.search(text_without_cost)
    code = code_match.group(0) if code_match else None

    name = _PERCENT_PREFIX_RE.sub("", text_without_cost)
    if code:
        name = name.replace(code, "", 1)
    name = _PORTION_RE.sub("", name)
    name = re.sub(r"\s+-\s+C51\b", "", name)
    name = re.sub(r"^\s*-\s*|\s*-\s*$", "", name)
    name = re.sub(r"\s{2,}", " ", name).strip()

    return name, portion, code


def normalize_planning_rows(rows: Iterable[Iterable[object]]) -> list[ImportedMenuItem]:
    rows = list(rows)
    if not rows:
        return []

    headers = [_normalize_header(value) for value in rows[0]]
    if not headers or headers[0] != "DIA":
        raise ValueError("layout de planejamento inválido: primeira coluna deve ser 'Dia'")

    imported: list[ImportedMenuItem] = []
    for row in rows[1:]:
        row = list(row)
        if not row or row[0] in (None, ""):
            continue
        try:
            day = int(str(row[0]).strip())
        except ValueError as exc:
            raise ValueError(f"dia inválido no cardápio: {row[0]!r}") from exc

        for index, header in enumerate(headers[1:], start=1):
            if not header or header in OPERATIONAL_COLUMNS:
                continue
            category = CATEGORY_MAP.get(header)
            if category is None:
                continue
            raw_value = row[index] if index < len(row) else None
            parsed = parse_planning_cell(raw_value)
            if parsed is None:
                continue
            name, portion, code = parsed
            if not name:
                continue
            imported.append(
                ImportedMenuItem(
                    day=day,
                    source_column=header,
                    category=category,
                    name=name,
                    portion=portion,
                    technical_sheet_code=code,
                    raw_value=str(raw_value),
                )
            )
    return imported


def read_planning_xlsx(path: str | Path, sheet_name: str = "Planejamento") -> list[ImportedMenuItem]:
    workbook = load_workbook(filename=path, read_only=True, data_only=True)
    if sheet_name not in workbook.sheetnames:
        raise ValueError(f"aba obrigatória não encontrada: {sheet_name}")
    sheet = workbook[sheet_name]
    return normalize_planning_rows(sheet.iter_rows(values_only=True))


def read_planning_csv(content: bytes | str, delimiter: str | None = None) -> list[ImportedMenuItem]:
    if isinstance(content, bytes):
        text = content.decode("utf-8-sig")
    else:
        text = content

    if delimiter is None:
        sample = text[:4096]
        try:
            delimiter = csv.Sniffer().sniff(sample, delimiters=";,\t,").delimiter
        except csv.Error:
            delimiter = ";"

    rows = csv.reader(io.StringIO(text), delimiter=delimiter)
    return normalize_planning_rows(rows)
