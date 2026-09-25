from __future__ import annotations

import csv
import io
import re
from dataclasses import asdict, dataclass
from decimal import Decimal, InvalidOperation
from pathlib import Path
from uuid import uuid4

from openpyxl import load_workbook
from sqlalchemy import text

from app.db import engine


@dataclass
class ImportedTechnicalSheet:
    code: str
    name: str
    category: str | None
    portion_quantity: Decimal | None
    portion_unit: str | None
    kcal: Decimal | None
    protein_g: Decimal | None
    carbs_g: Decimal | None
    fat_g: Decimal | None
    ingredients: list[str]
    allergens: list[dict[str, str]]


@dataclass
class StagedTechnicalSheetImport:
    preview_id: str
    file_name: str
    items: list[ImportedTechnicalSheet]
    warnings: list[str]


_STAGED: dict[str, StagedTechnicalSheetImport] = {}


def _key(value: object) -> str:
    text_value = str(value or '').strip().casefold()
    text_value = re.sub(r'[^a-z0-9áàãâéêíóôõúç]+', '_', text_value)
    return text_value.strip('_')


_ALIASES = {
    'code': {'codigo', 'código', 'codigo_tecnico', 'código_técnico', 'technical_sheet_code', 'ficha', 'ficha_tecnica'},
    'name': {'nome', 'preparacao', 'preparação', 'prato', 'name'},
    'category': {'categoria', 'category'},
    'portion_quantity': {'porcao', 'porção', 'quantidade_porção', 'quantidade_porcao', 'portion_quantity'},
    'portion_unit': {'unidade_porção', 'unidade_porcao', 'unidade', 'portion_unit'},
    'kcal': {'kcal', 'energia', 'calorias'},
    'protein_g': {'proteina', 'proteína', 'protein_g', 'proteina_g'},
    'carbs_g': {'carboidratos', 'carbo', 'carbs_g', 'carboidratos_g'},
    'fat_g': {'gorduras', 'gordura', 'fat_g', 'gordura_g'},
    'ingredients': {'ingredientes', 'ingredients'},
    'allergens': {'alergenicos', 'alergênicos', 'alergenos', 'alérgenos', 'allergens'},
}


def _canonical(header: object) -> str | None:
    normalized = _key(header)
    for canonical, aliases in _ALIASES.items():
        if normalized in {_key(alias) for alias in aliases}:
            return canonical
    return None


def _decimal(value: object) -> Decimal | None:
    if value is None or str(value).strip() == '':
        return None
    raw = str(value).strip().replace(',', '.')
    try:
        parsed = Decimal(raw)
    except InvalidOperation as exc:
        raise ValueError(f'valor numérico inválido: {value}') from exc
    if parsed < 0:
        raise ValueError(f'valor negativo não permitido: {value}')
    return parsed


def _split(value: object) -> list[str]:
    if value is None:
        return []
    return [part.strip() for part in re.split(r'[;|\n]+', str(value)) if part.strip()]


def _allergens(value: object) -> list[dict[str, str]]:
    result: list[dict[str, str]] = []
    for part in _split(value):
        if ':' in part:
            allergen, status = [x.strip() for x in part.split(':', 1)]
        else:
            allergen, status = part, 'contains'
        if status not in {'contains', 'may_contain', 'free_from'}:
            status = 'contains'
        if allergen:
            result.append({'allergen': allergen.casefold(), 'status': status})
    return result


def _parse_rows(rows: list[list[object]]) -> tuple[list[ImportedTechnicalSheet], list[str]]:
    if not rows:
        return [], ['Arquivo vazio.']
    headers = [_canonical(value) for value in rows[0]]
    if 'code' not in headers or 'name' not in headers:
        raise ValueError('A planilha precisa ter ao menos as colunas código e preparação/nome.')

    items: list[ImportedTechnicalSheet] = []
    warnings: list[str] = []
    seen: set[str] = set()
    for row_number, row in enumerate(rows[1:], start=2):
        data = {header: row[index] if index < len(row) else None for index, header in enumerate(headers) if header}
        code = str(data.get('code') or '').strip()
        name = str(data.get('name') or '').strip()
        if not code and not name:
            continue
        if not code or not name:
            warnings.append(f'Linha {row_number}: ignorada por faltar código ou nome.')
            continue
        if code in seen:
            warnings.append(f'Linha {row_number}: código duplicado {code}; última ocorrência será usada.')
            items = [item for item in items if item.code != code]
        seen.add(code)
        try:
            item = ImportedTechnicalSheet(
                code=code,
                name=name,
                category=(str(data.get('category')).strip() if data.get('category') not in (None, '') else None),
                portion_quantity=_decimal(data.get('portion_quantity')),
                portion_unit=(str(data.get('portion_unit')).strip() if data.get('portion_unit') not in (None, '') else None),
                kcal=_decimal(data.get('kcal')),
                protein_g=_decimal(data.get('protein_g')),
                carbs_g=_decimal(data.get('carbs_g')),
                fat_g=_decimal(data.get('fat_g')),
                ingredients=_split(data.get('ingredients')),
                allergens=_allergens(data.get('allergens')),
            )
        except ValueError as exc:
            warnings.append(f'Linha {row_number}: {exc}; registro ignorado.')
            continue
        items.append(item)
    return items, warnings


def read_import(file_name: str, content: bytes) -> tuple[list[ImportedTechnicalSheet], list[str]]:
    extension = Path(file_name).suffix.lower()
    if extension == '.csv':
        text_content = content.decode('utf-8-sig', errors='replace')
        sample = text_content[:4096]
        try:
            dialect = csv.Sniffer().sniff(sample, delimiters=',;\t')
        except csv.Error:
            dialect = csv.excel
            dialect.delimiter = ';'
        rows = [list(row) for row in csv.reader(io.StringIO(text_content), dialect)]
        return _parse_rows(rows)
    if extension == '.xlsx':
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        sheet = workbook.active
        rows = [list(row) for row in sheet.iter_rows(values_only=True)]
        workbook.close()
        return _parse_rows(rows)
    raise ValueError('Formato não suportado; envie .xlsx ou .csv.')


def stage_import(file_name: str, content: bytes) -> StagedTechnicalSheetImport:
    items, warnings = read_import(file_name, content)
    if not items:
        warnings.append('Nenhuma ficha válida foi encontrada.')
    staged = StagedTechnicalSheetImport(str(uuid4()), file_name, items, warnings)
    _STAGED[staged.preview_id] = staged
    return staged


def preview_payload(staged: StagedTechnicalSheetImport) -> dict:
    complete = sum(1 for item in staged.items if all(value is not None for value in (item.kcal, item.protein_g, item.carbs_g, item.fat_g)))
    return {
        'preview_id': staged.preview_id,
        'file_name': staged.file_name,
        'count': len(staged.items),
        'complete_count': complete,
        'incomplete_count': len(staged.items) - complete,
        'warnings': staged.warnings,
        'items': [asdict(item) for item in staged.items[:200]],
    }


def publish_import(preview_id: str) -> dict:
    staged = _STAGED.get(preview_id)
    if staged is None:
        raise LookupError('preview não encontrado ou expirado')
    created = 0
    updated = 0
    with engine.begin() as conn:
        for item in staged.items:
            exists = conn.execute(text('SELECT 1 FROM technical_sheets WHERE code = :code'), {'code': item.code}).scalar_one_or_none()
            conn.execute(text('''
                INSERT INTO technical_sheets
                    (code, name, category, portion_quantity, portion_unit, kcal, protein_g, carbs_g, fat_g)
                VALUES (:code,:name,:category,:portion_quantity,:portion_unit,:kcal,:protein_g,:carbs_g,:fat_g)
                ON CONFLICT (code) DO UPDATE SET
                    name=EXCLUDED.name, category=EXCLUDED.category,
                    portion_quantity=EXCLUDED.portion_quantity, portion_unit=EXCLUDED.portion_unit,
                    kcal=EXCLUDED.kcal, protein_g=EXCLUDED.protein_g, carbs_g=EXCLUDED.carbs_g,
                    fat_g=EXCLUDED.fat_g, updated_at=now()
            '''), {key: value for key, value in asdict(item).items() if key not in {'ingredients', 'allergens'}})
            conn.execute(text('DELETE FROM technical_sheet_ingredients WHERE technical_sheet_code=:code'), {'code': item.code})
            conn.execute(text('DELETE FROM technical_sheet_allergens WHERE technical_sheet_code=:code'), {'code': item.code})
            for position, ingredient in enumerate(item.ingredients, start=1):
                conn.execute(text('INSERT INTO technical_sheet_ingredients (technical_sheet_code,position,ingredient) VALUES (:code,:position,:ingredient)'), {'code': item.code, 'position': position, 'ingredient': ingredient})
            for allergen in item.allergens:
                conn.execute(text('INSERT INTO technical_sheet_allergens (technical_sheet_code,allergen,status) VALUES (:code,:allergen,:status)'), {'code': item.code, **allergen})
            if exists:
                updated += 1
            else:
                created += 1
    _STAGED.pop(preview_id, None)
    return {'status': 'published', 'created': created, 'updated': updated, 'count': created + updated}
